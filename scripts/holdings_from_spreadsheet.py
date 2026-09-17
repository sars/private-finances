#!/usr/bin/env python3
"""Turn the owner's asset spreadsheet into the JSON that
`node dist/src/holdings-import-cli.js` loads. Runs on the operator's machine;
the output names real holdings and must never be committed.

    python3 scripts/holdings_from_spreadsheet.py Assets.xlsx holdings.json
    python3 scripts/holdings_from_spreadsheet.py Assets.xlsx holdings.json \
        --assets CurrentAssets2 --rates Rates --mapping mapping.json

Layout this reads, and nothing else about the workbook:

* The assets sheet: row 1 carries a date in every other column starting at E
  (each date owns two columns, the native amount and a USD figure; only the
  amount is read). Row 2 is the header. Each later row with a name in A and a
  currency or symbol in B is a holding: C is the Invested flag and D the
  liquidity flag when they are booleans; a block without a C flag is the
  illiquid block, counted as invested, with D as the liquidity flag when
  present and not liquid otherwise.
* The rates sheet: rows of (date, symbol, units of the symbol per one USD).
  They are inverted to USD per unit here; USD itself is skipped.

An optional mapping file, kept outside Git, refines what the sheet cannot say:
    {"<sheet name>": {"name": "…", "kind": "bank", "group": "…",
                      "owner": "rodion", "archived": true, "skip": false}}

Requires openpyxl (`python3 -m pip install --user openpyxl`).
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime
from decimal import Context, Decimal, getcontext

getcontext().prec = 28

KINDS = {
    "cash", "bank", "broker", "crypto", "bond", "deposit", "fund",
    "real_estate", "business", "receivable", "other",
}


def as_date(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip()[:10], "%Y-%m-%d").date().isoformat()
        except ValueError:
            return None
    return None


def as_decimal(value) -> str | None:
    """A cell's number as a plain decimal string; errors and text are None."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        text = format(Decimal(str(value)).normalize(), "f")
        return text if text not in ("", "-0") else "0"
    if isinstance(value, str):
        stripped = value.strip().replace(",", "")
        try:
            return format(Decimal(stripped).normalize(), "f")
        except Exception:
            return None
    return None


def default_kind(denomination: str, liquid: bool) -> str:
    if denomination in ("BTC", "ETH"):
        return "crypto"
    if not liquid:
        return "other"
    if denomination in ("USD", "EUR", "UAH", "GBP"):
        return "bank"
    return "broker"


def read_assets(sheet, mapping: dict):
    dates: dict[int, str] = {}
    for column, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1)), start=1):
        day = as_date(cell.value)
        if day and column >= 5:
            dates[column] = day
    if not dates:
        raise SystemExit("no snapshot dates found in row 1")
    holdings, snapshots, seen = [], [], {}
    for row in sheet.iter_rows(min_row=3):
        if len(row) < 4:
            continue
        name, denom = row[0].value, row[1].value
        if not isinstance(name, str) or not name.strip():
            continue
        if not isinstance(denom, str) or not denom.strip():
            continue
        name = name.strip()
        denom = denom.strip().upper()
        c, d = row[2].value, row[3].value
        if isinstance(c, bool):
            invested, liquid = c, bool(d) if isinstance(d, bool) else True
        else:
            # The illiquid block carries no Invested flag; the sheet's own
            # "Not Invested" total leaves it out, so it counts as invested.
            invested, liquid = True, d if isinstance(d, bool) else False
        rules = mapping.get(name, {})
        if rules.get("skip"):
            continue
        final_name = rules.get("name", name)
        if final_name in seen:
            seen[final_name] += 1
            final_name = f"{final_name} ({seen[final_name]})"
        else:
            seen[final_name] = 1
        kind = rules.get("kind", default_kind(denom, liquid))
        if kind not in KINDS:
            raise SystemExit(f"unknown kind {kind!r} in mapping")
        holding = {
            "name": final_name,
            "denomination": denom,
            "kind": kind,
            "invested": bool(rules.get("invested", invested)),
            "liquid": bool(rules.get("liquid", liquid)),
        }
        for key in ("group", "owner", "note"):
            if rules.get(key) is not None:
                holding[key] = rules[key]
        if rules.get("archived"):
            holding["archived"] = True
        if "sortOrder" in rules:
            holding["sortOrder"] = int(rules["sortOrder"])
        holdings.append(holding)
        for column, cell in enumerate(row, start=1):
            day = dates.get(column)
            if day is None:
                continue
            quantity = as_decimal(cell.value)
            if quantity is None:
                continue
            snapshots.append({"name": final_name, "asOf": day, "quantity": quantity})
    return holdings, snapshots


def read_rates(sheet):
    prices = []
    for row in sheet.iter_rows(min_row=1, values_only=True):
        if len(row) < 3:
            continue
        day, symbol, per_usd = as_date(row[0]), row[1], as_decimal(row[2])
        if not day or not isinstance(symbol, str) or per_usd is None:
            continue
        symbol = symbol.strip().upper()
        if symbol == "USD":
            continue
        units = Decimal(per_usd)
        if units <= 0:
            continue
        # Eight significant digits: the sheet's rates are floats and carry noise.
        usd_per_unit = Context(prec=8).divide(Decimal(1), units)
        prices.append({
            "symbol": symbol,
            "asOf": day,
            "usdPerUnit": format(usd_per_unit.normalize(), "f"),
        })
    return prices


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("workbook")
    parser.add_argument("output")
    parser.add_argument("--assets", default="CurrentAssets2")
    parser.add_argument("--rates", default="Rates")
    parser.add_argument("--mapping")
    args = parser.parse_args()
    try:
        import openpyxl
    except ImportError:
        print("openpyxl is required: python3 -m pip install --user openpyxl", file=sys.stderr)
        return 2
    mapping = {}
    if args.mapping:
        with open(args.mapping, encoding="utf-8") as handle:
            mapping = json.load(handle)
    workbook = openpyxl.load_workbook(args.workbook, data_only=True, read_only=True)
    holdings, snapshots = read_assets(workbook[args.assets], mapping)
    prices = read_rates(workbook[args.rates]) if args.rates in workbook.sheetnames else []
    document = {"holdings": holdings, "snapshots": snapshots, "prices": prices}
    with open(args.output, "w", encoding="utf-8") as handle:
        json.dump(document, handle, ensure_ascii=False, indent=1)
    # Counts only: the names and figures stay in the file.
    print(json.dumps({
        "holdings": len(holdings),
        "snapshots": len(snapshots),
        "prices": len(prices),
        "dates": len({s["asOf"] for s in snapshots}),
    }))
    return 0


if __name__ == "__main__":
    sys.exit(main())
